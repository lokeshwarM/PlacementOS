import io
import pytest
import httpx
import openpyxl
import docx
from pypdf import PdfWriter
from app.main import app
from app.document_processor import process_document
from app.doc_classifier import classify_document
from app.excel_parser import parse_excel_document
from app.pdf_parser import parse_pdf_document
from app.docx_parser import parse_docx_document

# =============================================================================
# 1. Excel Parsing Tests
# =============================================================================

def test_excel_parsing_standard():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Shortlisted"
    ws.append(["S.No", "Registration Number", "Candidate Name", "Applied Role"])
    ws.append([1, "21BCE1001", "Alice Smith", "Software Engineer"])
    ws.append([2, "21BCE1002", "Bob Johnson", "Data Scientist"])
    ws.append([3, "", "", ""]) # Blank row
    ws.append([4, "21BCE1001", "Alice Smith", "Software Engineer"]) # Duplicate

    buf = io.BytesIO()
    wb.save(buf)
    file_bytes = buf.getvalue()

    candidates = parse_excel_document(file_bytes)
    assert len(candidates) == 2
    assert candidates[0].registration_number == "21BCE1001"
    assert candidates[0].name == "Alice Smith"
    assert candidates[0].role == "Software Engineer"
    assert candidates[1].registration_number == "21BCE1002"
    assert candidates[1].name == "Bob Johnson"

def test_excel_parsing_offset_header_and_neopat():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Round2"
    # Offset metadata in first 3 rows
    ws.append(["Company: Acme Corp"])
    ws.append(["Date: 2026-09-01"])
    ws.append([])
    # Header on row 4
    ws.append(["NeoPAT ID", "Reg. No.", "Student Name"])
    ws.append(["NP1001", "21BIT2001", "Charlie Brown"])
    ws.append(["NP1002", "21BIT2002", "Diana Prince"])

    buf = io.BytesIO()
    wb.save(buf)
    file_bytes = buf.getvalue()

    candidates = parse_excel_document(file_bytes)
    assert len(candidates) == 2
    assert candidates[0].neopat_id == "NP1001"
    assert candidates[0].registration_number == "21BIT2001"
    assert candidates[0].name == "Charlie Brown"

def test_excel_multiple_sheets():
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "CSE"
    ws1.append(["Reg No", "Name"])
    ws1.append(["21BCE3001", "Edward Norton"])

    ws2 = wb.create_sheet(title="ECE")
    ws2.append(["Reg No", "Name"])
    ws2.append(["21BEC4001", "Fiona Gallagher"])

    buf = io.BytesIO()
    wb.save(buf)
    file_bytes = buf.getvalue()

    candidates = parse_excel_document(file_bytes)
    assert len(candidates) == 2
    assert {c.registration_number for c in candidates} == {"21BCE3001", "21BEC4001"}

# =============================================================================
# 2. PDF Parsing Tests
# =============================================================================

def test_pdf_parsing_scanned_requires_ocr():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    file_bytes = buf.getvalue()

    candidates, ocr_required, text = parse_pdf_document(file_bytes)
    assert ocr_required is True
    assert len(candidates) == 0

# =============================================================================
# 3. DOCX Parsing Tests
# =============================================================================

def test_docx_table_parsing():
    doc = docx.Document()
    doc.add_heading("Shortlisted Candidates", level=1)
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Roll Number"
    hdr_cells[1].text = "Candidate Name"
    hdr_cells[2].text = "Job Role"

    row_cells = table.add_row().cells
    row_cells[0].text = "21BCE5001"
    row_cells[1].text = "George Clark"
    row_cells[2].text = "Full Stack Engineer"

    buf = io.BytesIO()
    doc.save(buf)
    file_bytes = buf.getvalue()

    candidates = parse_docx_document(file_bytes)
    assert len(candidates) == 1
    assert candidates[0].registration_number == "21BCE5001"
    assert candidates[0].name == "George Clark"
    assert candidates[0].role == "Full Stack Engineer"

def test_docx_paragraph_parsing():
    doc = docx.Document()
    doc.add_heading("Selected Students", level=1)
    doc.add_paragraph("1. 21BCE6001 - Hannah Montana")
    doc.add_paragraph("2. 21BCE6002 - Ian Malcolm")

    buf = io.BytesIO()
    doc.save(buf)
    file_bytes = buf.getvalue()

    candidates = parse_docx_document(file_bytes)
    assert len(candidates) == 2
    assert {c.registration_number for c in candidates} == {"21BCE6001", "21BCE6002"}

# =============================================================================
# 4. Document Classifier Tests
# =============================================================================

def test_classifier_intent():
    assert classify_document("Google_Shortlisted_Candidates.xlsx")[0] == "SHORTLIST"
    assert classify_document("Microsoft_Interview_Shortlist_Round1.pdf")[0] == "SHORTLIST"
    assert classify_document("Amazon_Job_Description_2027.pdf")[0] == "JOB_DESCRIPTION"
    assert classify_document("Eligible_Students_List.xlsx")[0] == "ELIGIBILITY_DOCUMENT"
    assert classify_document("Recruitment_Schedule.docx")[0] == "RECRUITMENT_PROCESS"
    assert classify_document("Campus_Map.pdf")[0] == "OTHER"

# =============================================================================
# 5. FastAPI Endpoints & Full Processing Tests
# =============================================================================

def test_process_document_excel_endpoint():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Registration Number", "Name"])
    ws.append(["21BCE7001", "Julia Roberts"])

    buf = io.BytesIO()
    wb.save(buf)
    file_bytes = buf.getvalue()

    res = process_document(101, "Shortlist_Round1.xlsx", file_bytes)
    assert res.classification == "SHORTLIST"
    assert res.document_type == "XLSX"
    assert len(res.candidates) == 1
    assert res.candidates[0].registration_number == "21BCE7001"
    assert res.candidates[0].name == "Julia Roberts"

@pytest.mark.asyncio
async def test_fastapi_document_endpoints():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Reg No", "Candidate Name"])
    ws.append(["21BCE8001", "Kevin Bacon"])

    buf = io.BytesIO()
    wb.save(buf)
    file_bytes = buf.getvalue()

    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as ac:
        response = await ac.post(
            "/api/v1/documents/process",
            files={"file": ("Google_Shortlist.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"attachment_id": 202}
        )
        assert response.status_code == 200
        data = response.json()
        assert data["attachmentId"] == 202
        assert data["classification"] == "SHORTLIST"
        assert len(data["candidates"]) == 1
        assert data["candidates"][0]["registrationNumber"] == "21BCE8001"
