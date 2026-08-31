import io
import re
from typing import List
import docx
from app.models import DocumentCandidate
from app.excel_parser import find_header_row, matches_any, REG_NO_PATTERNS, NEOPAT_PATTERNS, NAME_PATTERNS, ROLE_PATTERNS

REG_NO_REGEX = re.compile(r"\b(\d{2}[A-Za-z]{3}\d{4,5})\b")
NEOPAT_REGEX = re.compile(r"\b(NP\d{3,8}|NEO\d{3,8})\b", re.IGNORECASE)

def parse_docx_document(file_bytes: bytes) -> List[DocumentCandidate]:
    """
    Parses a DOCX document, extracting candidate rows from tables and bulleted/formatted paragraphs.
    """
    candidates: List[DocumentCandidate] = []
    seen_keys = set()

    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        return candidates

    # 1. Parse Tables
    for t_idx, table in enumerate(doc.tables):
        table_rows = []
        for row in table.rows:
            table_rows.append(tuple(cell.text.strip() for cell in row.cells))

        if not table_rows:
            continue

        header_idx, col_map = find_header_row(table_rows)
        if header_idx is None:
            header_idx = 0
            # Detect by regex if needed
            col_map = {}
            for r in table_rows:
                for c_idx, text in enumerate(r):
                    if REG_NO_REGEX.search(text) and "reg_no" not in col_map:
                        col_map["reg_no"] = c_idx

        for r_idx in range(header_idx + 1, len(table_rows)):
            row = table_rows[r_idx]
            if not any(row):
                continue

            reg_no = row[col_map["reg_no"]] if "reg_no" in col_map and col_map["reg_no"] < len(row) else None
            neopat = row[col_map["neopat"]] if "neopat" in col_map and col_map["neopat"] < len(row) else None
            name = row[col_map["name"]] if "name" in col_map and col_map["name"] < len(row) else None
            role = row[col_map["role"]] if "role" in col_map and col_map["role"] < len(row) else None

            # If reg_no / neopat were not cleanly separated into columns, search cell contents
            if not reg_no:
                for cell_str in row:
                    m = REG_NO_REGEX.search(cell_str)
                    if m:
                        reg_no = m.group(1).upper()
                        break

            if not neopat:
                for cell_str in row:
                    m = NEOPAT_REGEX.search(cell_str)
                    if m:
                        neopat = m.group(1).upper()
                        break

            if not reg_no and not neopat and not name:
                continue

            key = (reg_no or "", neopat or "", name or "")
            if key in seen_keys:
                continue
            seen_keys.add(key)

            evidence = f"DOCX Table {t_idx + 1}, Row {r_idx + 1}: {' | '.join(row[:4])}"
            candidates.append(DocumentCandidate(
                registration_number=reg_no,
                neopat_id=neopat,
                name=name,
                role=role,
                evidence=evidence
            ))

    # 2. Parse Paragraphs (for bulleted or line-by-line shortlists)
    for p_idx, p in enumerate(doc.paragraphs):
        p_text = p.text.strip()
        if not p_text:
            continue

        reg_match = REG_NO_REGEX.search(p_text)
        neo_match = NEOPAT_REGEX.search(p_text)

        if reg_match or neo_match:
            reg_no = reg_match.group(1).upper() if reg_match else None
            neopat = neo_match.group(1).upper() if neo_match else None

            rem = p_text
            if reg_no: rem = rem.replace(reg_no, " ")
            if neopat: rem = rem.replace(neopat, " ")
            rem = re.sub(r"^[0-9]+[\.\-\)\s]+", "", rem).strip() # strip leading numbering
            name = rem if len(rem) > 2 else None

            key = (reg_no or "", neopat or "", name or "")
            if key not in seen_keys:
                seen_keys.add(key)
                candidates.append(DocumentCandidate(
                    registration_number=reg_no,
                    neopat_id=neopat,
                    name=name,
                    role=None,
                    evidence=f"DOCX Paragraph {p_idx + 1}: {p_text[:100]}"
                ))

    return candidates
